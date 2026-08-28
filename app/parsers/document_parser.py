from dataclasses import dataclass
from pathlib import Path
import re

import fitz
from docx import Document as DocxDocument


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


@dataclass(slots=True)
class ParsedPage:
    text: str
    page_number: int | None = None
    section: str | None = None


@dataclass(slots=True)
class ParsedDocument:
    text: str
    chunks: list[str]
    pages: list[ParsedPage]


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_pdf(path: Path) -> list[ParsedPage]:
    pages: list[ParsedPage] = []
    with fitz.open(path) as pdf:
        for index, page in enumerate(pdf, start=1):
            text = normalize_text(page.get_text())
            if text:
                pages.append(ParsedPage(text=text, page_number=index))
    return pages


def parse_docx(path: Path) -> list[ParsedPage]:
    document = DocxDocument(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return [ParsedPage(text=normalize_text("\n".join(paragraphs)))] if paragraphs else []


def parse_plain_text(path: Path) -> list[ParsedPage]:
    text = normalize_text(read_text_file(path))
    return [ParsedPage(text=text)] if text else []


def split_text(text: str, chunk_size: int = 1000, overlap: int = 150) -> list[str]:
    text = normalize_text(text)
    if not text:
        return []
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")

    chunks: list[str] = []
    start = 0
    text_length = len(text)
    while start < text_length:
        end = min(start + chunk_size, text_length)
        if end < text_length:
            boundary_candidates = [
                text.rfind("\n\n", start, end),
                text.rfind("\n", start, end),
                text.rfind("。", start, end),
                text.rfind(" ", start, end),
            ]
            boundary = max(boundary_candidates)
            if boundary > start + chunk_size // 2:
                end = boundary + (2 if text[boundary:boundary + 2] == "\n\n" else 1)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= text_length:
            break
        start = max(end - overlap, start + 1)
    return chunks


def parse_document(path: str | Path) -> ParsedDocument:
    file_path = Path(path)
    extension = file_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型：{extension or 'unknown'}")

    if extension == ".pdf":
        pages = parse_pdf(file_path)
    elif extension == ".docx":
        pages = parse_docx(file_path)
    else:
        pages = parse_plain_text(file_path)

    text = normalize_text("\n\n".join(page.text for page in pages))
    return ParsedDocument(text=text, chunks=split_text(text), pages=pages)
