import asyncio
import os
import uuid

import fitz
import pytest
import pytest_asyncio
from docx import Document as DocxDocument
from httpx import ASGITransport, AsyncClient

from app.db.init_db import init_db
from app.main import app


pytestmark = pytest.mark.integration


@pytest_asyncio.fixture
async def client() -> AsyncClient:
    if os.getenv("RUN_DB_TESTS") != "1":
        pytest.skip("Set RUN_DB_TESTS=1 to run PostgreSQL integration tests")
    await init_db()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as test_client:
        yield test_client


def make_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Employee Expense Policy\nTravel and hotel costs require valid receipts.")
    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("请假制度", level=1)
    document.add_paragraph("员工请假应提前提交申请，并等待审批。")
    buffer = __import__("io").BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def wait_for_document(client: AsyncClient, document_id: str, token: str) -> dict:
    headers = {"Authorization": f"Bearer {token}"}
    for _ in range(40):
        response = await client.get(f"/api/v1/documents/{document_id}", headers=headers)
        assert response.status_code == 200, response.text
        document = response.json()
        if document["status"] in {"completed", "failed"}:
            return document
        await asyncio.sleep(0.25)
    raise AssertionError("document processing timed out")


async def register_and_login(client: AsyncClient) -> str:
    email = f"documents-{uuid.uuid4().hex}@example.com"
    registered = await client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "StrongPass123!"},
    )
    assert registered.status_code == 201, registered.text
    logged_in = await client.post(
        "/api/v1/auth/login",
        data={"username": email, "password": "StrongPass123!"},
    )
    assert logged_in.status_code == 200, logged_in.text
    return logged_in.json()["access_token"]


@pytest.mark.asyncio
async def test_upload_and_process_supported_documents(client: AsyncClient) -> None:
    token = await register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    created = await client.post(
        "/api/v1/knowledge-bases",
        headers=headers,
        json={"name": "第三天文档测试"},
    )
    assert created.status_code == 201, created.text
    knowledge_base_id = created.json()["id"]

    files = [
        ("policy.txt", "员工考勤制度\n迟到需要按制度处理。".encode("utf-8"), "text/plain"),
        ("policy.md", "# 报销制度\n\n交通费需要保留发票。".encode("utf-8"), "text/markdown"),
        ("policy.pdf", make_pdf(), "application/pdf"),
        ("policy.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ]
    document_ids = []
    for filename, content, content_type in files:
        response = await client.post(
            f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
            headers=headers,
            files={"file": (filename, content, content_type)},
        )
        assert response.status_code == 202, response.text
        body = response.json()
        document_ids.append(body["document"]["id"])
        assert body["document"]["status"] in {"queued", "processing"}

    processed = [await wait_for_document(client, document_id, token) for document_id in document_ids]
    assert all(item["status"] == "completed" for item in processed), processed
    assert all(item["chunk_count"] >= 1 for item in processed), processed

    listed = await client.get(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents", headers=headers
    )
    assert listed.status_code == 200
    assert len(listed.json()) >= 4


@pytest.mark.asyncio
async def test_document_upload_validation(client: AsyncClient) -> None:
    token = await register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    created = await client.post(
        "/api/v1/knowledge-bases",
        headers=headers,
        json={"name": "第三天校验测试"},
    )
    knowledge_base_id = created.json()["id"]

    unsupported = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": ("data.xlsx", b"not supported", "application/octet-stream")},
    )
    assert unsupported.status_code == 415

    empty = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": ("empty.txt", b"", "text/plain")},
    )
    assert empty.status_code == 400
@pytest.mark.asyncio
async def test_hybrid_search_returns_indexed_chunks(client: AsyncClient) -> None:
    token = await register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    created = await client.post(
        "/api/v1/knowledge-bases",
        headers=headers,
        json={"name": "第四天检索测试"},
    )
    assert created.status_code == 201, created.text
    knowledge_base_id = created.json()["id"]

    response = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={
            "file": (
                "expense-policy.txt",
                "Expense reimbursement policy: transportation and hotel costs require valid receipts.".encode("utf-8"),
                "text/plain",
            )
        },
    )
    assert response.status_code == 202, response.text
    document = await wait_for_document(client, response.json()["document"]["id"], token)
    assert document["status"] == "completed", document

    search = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/search",
        headers=headers,
        json={"query": "hotel costs receipts", "top_k": 5},
    )
    assert search.status_code == 200, search.text
    body = search.json()
    assert body["embedding_provider"] in {"hash-fallback", "openai-compatible"}
    assert body["result_count"] >= 1
    assert body["results"][0]["document_name"] == "expense-policy.txt"
    assert body["results"][0]["vector_score"] is not None
    assert body["results"][0]["hybrid_score"] > 0

