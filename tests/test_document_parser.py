from io import BytesIO

import fitz
from docx import Document as DocxDocument

from app.parsers.document_parser import normalize_text, parse_document, split_text


def write_pdf(path):
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Employee Expense Policy\nTravel and hotel costs require valid receipts.")
    document.save(path)
    document.close()


def write_docx(path):
    document = DocxDocument()
    document.add_heading("请假制度", level=1)
    document.add_paragraph("员工请假应提前提交申请，并等待审批。")
    document.save(path)


def test_normalize_and_split():
    text = normalize_text("  第一行\r\n\r\n\r\n 第二行  ")
    assert text == "第一行\n\n 第二行"
    chunks = split_text("这是第一段。\n\n这是第二段。", chunk_size=8, overlap=2)
    assert chunks
    assert all(chunk.strip() for chunk in chunks)


def test_parse_txt_and_markdown(tmp_path):
    txt = tmp_path / "policy.txt"
    txt.write_bytes("员工考勤制度\n迟到需要按制度处理。".encode("utf-8"))
    parsed_txt = parse_document(txt)
    assert "员工考勤制度" in parsed_txt.text
    assert parsed_txt.chunks

    markdown = tmp_path / "policy.md"
    markdown.write_text("# 报销制度\n\n交通费需要保留发票。", encoding="utf-8")
    parsed_md = parse_document(markdown)
    assert "报销制度" in parsed_md.text


def test_parse_pdf_and_docx(tmp_path):
    pdf = tmp_path / "policy.pdf"
    write_pdf(pdf)
    parsed_pdf = parse_document(pdf)
    assert "Employee Expense Policy" in parsed_pdf.text
    assert parsed_pdf.pages[0].page_number == 1

    docx = tmp_path / "policy.docx"
    write_docx(docx)
    parsed_docx = parse_document(docx)
    assert "请假制度" in parsed_docx.text


def test_unsupported_file(tmp_path):
    unsupported = tmp_path / "policy.xlsx"
    unsupported.write_bytes(b"not supported")
    try:
        parse_document(unsupported)
    except ValueError as exc:
        assert "不支持" in str(exc)
    else:
        raise AssertionError("unsupported extension should raise ValueError")

